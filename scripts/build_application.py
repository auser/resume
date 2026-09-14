#!/usr/bin/env python3
"""Build the targeted resume and review pack from application-profile.json.

Requires Python 3.10+ and python-docx. PDF output additionally requires
LibreOffice (soffice/libreoffice). No network requests or external services.
The original resume.json and Typst templates are never modified.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def load_profile(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    required = {"name", "target", "summary", "contact_lines", "skills", "experience", "education", "cover_letter_paragraphs"}
    missing = required - data.keys()
    if missing:
        raise ValueError(f"Missing profile fields: {', '.join(sorted(missing))}")
    words = len(" ".join(data["cover_letter_paragraphs"]).split())
    if words != 120:
        raise ValueError(f"Cover letter must contain 120 whitespace-delimited words; found {words}.")
    for job in data["experience"]:
        ids = [b["source_bullet"] for b in job["highlights"]]
        if len(ids) != len(set(ids)):
            raise ValueError(f"Duplicate source bullet in {job['company']}")
        for b in job["highlights"]:
            if not b["text"].strip() or not isinstance(b["include"], bool):
                raise ValueError(f"Invalid bullet in {job['company']}")
    return data


def selected(job: dict) -> list[dict]:
    return [b for b in job["highlights"] if b["include"]]


def heading(data: dict) -> str:
    return f"{data['target']['title']} | {data['target']['focus']}"


def resume_markdown(data: dict) -> str:
    lines = [f"# {data['name']}", "", heading(data), "", *data["contact_lines"], "", "## Summary", "", data["summary"], "", "## Skills", ""]
    lines += [f"**{s['label']}:** {s['text']}" for s in data["skills"]]
    for earlier, title in [(False, "Experience"), (True, "Earlier Experience")]:
        lines += ["", f"## {title}", ""]
        for job in data["experience"]:
            if job["earlier"] != earlier or not selected(job):
                continue
            lines += [f"### {job['title']} | {job['company']}", job["period"], ""]
            lines += [f"- {b['text']}" for b in selected(job)]
            lines.append("")
    edu = data["education"]
    lines += ["## Education", "", f"{edu['institution']} | {edu['period']}", "; ".join(edu["degrees"]), ""]
    return "\n".join(lines)


def bullet_bank(data: dict) -> str:
    src = data["source"]
    total = sum(len(j["highlights"]) for j in data["experience"])
    kept = sum(len(selected(j)) for j in data["experience"])
    lines = ["# Complete outcome-led bullet bank", "", f"{total} experience bullets rewritten; {kept} selected for the application resume.", "", f"Source: `{src['repository']}/{src['path']}` at commit `{src['commit']}`. Each number below maps to the original bullet within that role.", "", "Source-reported claims are not independently audited. Unselected bullets are a review bank, not all approved for submission. No estimated results have been invented.", ""]
    for job in data["experience"]:
        lines += [f"## {job['company']} | {job['title']} | {job['period']}", ""]
        for b in job["highlights"]:
            status = "SELECTED" if b["include"] else "OMITTED FROM APPLICATION"
            lines += [f"### Original bullet {b['source_bullet']} — {status}", "", b["text"], ""]
            if b.get("review_note"):
                lines += [f"Review: {b['review_note']}", ""]
    return "\n".join(lines)


def cover_letter(data: dict) -> str:
    return "\n\n".join(data["cover_letter_paragraphs"]) + "\n"


def build_docx(data: dict, path: Path) -> None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.style import WD_STYLE_TYPE
    except ImportError as exc:
        raise RuntimeError("Install the document dependency: python -m pip install python-docx") from exc
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.6)
    sec.left_margin, sec.right_margin = Inches(0.7), Inches(0.7)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(4)
    for key in ["Heading 1", "Heading 2"]:
        styles[key].font.name = "Arial"
        styles[key].font.size = Pt(10.5)
        styles[key].font.bold = True
        styles[key].font.color.rgb = RGBColor(0, 0, 0)
        styles[key].paragraph_format.space_before = Pt(9 if key == "Heading 1" else 6)
        styles[key].paragraph_format.space_after = Pt(3)
        styles[key].paragraph_format.keep_with_next = True
    lb = styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    lb.base_style = normal
    lb.font.name = "Arial"
    lb.font.size = Pt(10.5)
    lb.paragraph_format.left_indent = Inches(0.14)
    lb.paragraph_format.first_line_indent = Inches(-0.14)
    lb.paragraph_format.space_after = Pt(3)
    lb.paragraph_format.line_spacing = 1.05
    lb.paragraph_format.widow_control = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(data["name"])
    r.bold, r.font.size = True, Pt(22)
    p = doc.add_paragraph(heading(data))
    p.paragraph_format.space_after = Pt(4)
    for text in data["contact_lines"]:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.size = Pt(9.5)
    doc.add_heading("SUMMARY", level=1)
    doc.add_paragraph(data["summary"])
    doc.add_heading("SKILLS", level=1)
    for skill in data["skills"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.add_run(skill["label"] + ": ").bold = True
        p.add_run(skill["text"])
    doc.add_heading("EXPERIENCE", level=1)
    earlier_started = False
    for job in data["experience"]:
        if not selected(job):
            continue
        if job["company"] == "Amazon Web Services":
            doc.add_page_break()
            p = doc.add_heading("EXPERIENCE (CONTINUED)", level=1)
            p.paragraph_format.space_before = Pt(0)
        if job["earlier"] and not earlier_started:
            doc.add_heading("EARLIER EXPERIENCE", level=1)
            earlier_started = True
        doc.add_heading(f"{job['title']} | {job['company']}", level=2)
        p = doc.add_paragraph(job["period"])
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.font.size = Pt(9.5)
        for b in selected(job):
            p = doc.add_paragraph(style="Resume Bullet")
            p.add_run("- ")
            p.paragraph_format.keep_together = True
            # Emphasize only an actual numeric prefix; no inserted claims.
            match = re.match(r"(\$?\d[\d,.]*\+?%?)(.*)", b["text"])
            if match:
                p.add_run(match.group(1)).bold = True
                p.add_run(match.group(2))
            else:
                p.add_run(b["text"])
    edu = data["education"]
    doc.add_heading("EDUCATION", level=1)
    p = doc.add_paragraph()
    p.add_run(edu["institution"]).bold = True
    p.add_run(" | " + edu["period"])
    p.paragraph_format.keep_with_next = True
    doc.add_paragraph("; ".join(edu["degrees"]))
    doc.core_properties.author = data["name"]
    doc.core_properties.title = "Ari Lerner - Platform Infrastructure and AI Security"
    doc.core_properties.subject = "Application resume; target role is provisional"
    doc.save(path)


def convert_pdf(docx: Path, out_dir: Path) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if executable is None and mac.is_file():
        executable = str(mac)
    if executable is None:
        raise RuntimeError("PDF export requires LibreOffice. The DOCX and text files have still been created.")
    # Fresh profile and output directory avoid locking and stale-output success.
    with tempfile.TemporaryDirectory(prefix="resume-pdf-") as temp:
        temp_root = Path(temp)
        pdf_dir = temp_root / "pdf"
        pdf_dir.mkdir()
        command = [executable, f"-env:UserInstallation={(temp_root/'profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(pdf_dir), str(docx.resolve())]
        result = subprocess.run(command, capture_output=True, text=True, timeout=90)
        produced = pdf_dir / (docx.stem + ".pdf")
        if result.returncode != 0 or not produced.is_file() or produced.stat().st_size == 0:
            raise RuntimeError(f"LibreOffice export failed: {result.stdout}\n{result.stderr}")
        destination = out_dir / produced.name
        shutil.copy2(produced, destination)
        return destination


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--profile", type=Path, default=ROOT / "application-profile.json")
    parser.add_argument("--out-dir", type=Path, default=ROOT / "build" / "application")
    parser.add_argument("--pdf", action="store_true", help="Also export PDF using LibreOffice")
    args = parser.parse_args()
    try:
        data = load_profile(args.profile)
        out = args.out_dir.resolve()
        out.mkdir(parents=True, exist_ok=True)
        md = resume_markdown(data)
        (out / "ari-lerner-resume-application.md").write_text(md, encoding="utf-8")
        plain = re.sub(r"(?m)^#{1,3} ", "", md).replace("**", "")
        (out / "ari-lerner-resume-application.txt").write_text(plain, encoding="utf-8")
        (out / "cover-letter-120-words.txt").write_text(cover_letter(data), encoding="utf-8")
        (out / "complete-bullet-rewrite.md").write_text(bullet_bank(data), encoding="utf-8")
        docx = out / "ari-lerner-resume-application.docx"
        build_docx(data, docx)
        if args.pdf:
            convert_pdf(docx, out)
        print(f"Created application files in {out}")
        print(f"Rewritten bullets: {sum(len(j['highlights']) for j in data['experience'])}; selected: {sum(len(selected(j)) for j in data['experience'])}; cover letter: 120 words")
        return 0
    except (OSError, ValueError, KeyError, TypeError, RuntimeError, subprocess.SubprocessError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
